# Ultralytics 🚀 AGPL-3.0 License - https://ultralytics.com/license

import os
import sys
import docx
import logging

# Ensure project root is in path
project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if project_root not in sys.path:
    sys.path.insert(0, project_root)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

DOCX_PATHS = [
    "/mnt/hdd/ultralytics/Computer_Vision_CCTV_Tracking_Project_Report.docx",
    "/mnt/hdd/ultralytics/project/Computer_Vision_CCTV_Tracking_Project_Report.docx"
]

def update_toc():
    for doc_path in DOCX_PATHS:
        if not os.path.exists(doc_path):
            logger.warning(f"File not found: {doc_path}")
            continue

        logger.info(f"Updating Table of Contents in: {doc_path}")
        doc = docx.Document(doc_path)
        t = doc.tables[0]
        
        # Row 5 (Chapter II)
        t.cell(5, 2).text = "12–25"
        # Row 6 (Chapter III)
        t.cell(6, 2).text = "26–30"
        # Row 7 (Chapter IV)
        t.cell(7, 2).text = "31–38"
        # Row 8 (Chapter V)
        t.cell(8, 2).text = "39–45"
        # Row 9 (Chapter VI)
        t.cell(9, 2).text = "46–51"
        # Row 10 (Chapter VII)
        t.cell(10, 2).text = "52–57"
        # Row 11 (Bibliography)
        t.cell(11, 2).text = "58–60"
        # Row 12 (Annexures)
        t.cell(12, 2).text = "61–67"

        doc.save(doc_path)
        logger.info(f"Saved successfully: {doc_path}")

if __name__ == "__main__":
    update_toc()
