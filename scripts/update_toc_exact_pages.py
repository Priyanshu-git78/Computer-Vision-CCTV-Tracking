# Ultralytics 🚀 AGPL-3.0 License - https://ultralytics.com/license

import os
import sys
import docx
import logging

# Ensure project root is in path
project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if project_root not in sys.path:
    sys.path.insert(0, project_root)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

DOCX_PATHS = [
    "/mnt/hdd/ultralytics/Computer_Vision_CCTV_Tracking_Project_Report.docx",
    "/mnt/hdd/ultralytics/project/Computer_Vision_CCTV_Tracking_Project_Report.docx"
]

def update_toc_exact():
    for doc_path in DOCX_PATHS:
        if not os.path.exists(doc_path):
            logger.warning(f"File not found: {doc_path}")
            continue

        logger.info(f"Updating Table of Contents with exact page numbers in: {doc_path}")
        doc = docx.Document(doc_path)
        t = doc.tables[0]
        
        # Row 1 (Table of Contents)
        t.cell(1, 2).text = "6"
        # Row 2 (List of Abbreviations)
        t.cell(2, 2).text = "7–8"
        # Row 3 (Abstract)
        t.cell(3, 2).text = "9"
        # Row 4 (Chapter I)
        t.cell(4, 2).text = "10–11"
        # Row 5 (Chapter II)
        t.cell(5, 2).text = "12–13"
        # Row 6 (Chapter III)
        t.cell(6, 2).text = "14–15"
        # Row 7 (Chapter IV)
        t.cell(7, 2).text = "16–20"
        # Row 8 (Chapter V)
        t.cell(8, 2).text = "21–24"
        # Row 9 (Chapter VI)
        t.cell(9, 2).text = "25–26"
        # Row 10 (Chapter VII)
        t.cell(10, 2).text = "27–28"
        # Row 11 (Bibliography)
        t.cell(11, 2).text = "29"
        # Row 12 (Annexures)
        t.cell(12, 2).text = "30–33"

        doc.save(doc_path)
        logger.info(f"Saved successfully: {doc_path}")

if __name__ == "__main__":
    update_toc_exact()
